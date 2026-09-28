"""L2（VLAN 进程内）：发现三相（§4.1）+ 三件产物 + §3.3 标记 + 地址漂移重发现。

含「脏名设备」（空格/中文/超长/重名/空名）与 RPM 不支持设备的兜底枚举路径。
"""

from __future__ import annotations

import json
import uuid
from typing import ClassVar

import pytest
from bacpypes3.local.analog import AnalogValueObject
from bacpypes3.vlan import VirtualNetwork

from tests.sim.device import SimDevice
from thermio_gateway.bacnet.discover import (
    Discoverer,
    stage_pointmap,
    write_export_xlsx,
    write_report,
)
from thermio_gateway.bacnet.stack import BacnetStack, build_virtual_application
from thermio_gateway.db import open_db
from thermio_gateway.pointmap import adopt_staging, load_pointmap


class DirtySimDevice(SimDevice):
    """现场质量不可控点名的设备（§3.3 各分支触发器）。

    空名分支用全空白名模拟（BACnet objectName 最小长度 1，全空白经
    trim 后为空 → generated 分支，与真实现场「名未配置」同效果）。
    """

    DIRTY_OBJECTS: ClassVar[dict] = {
        "CHWS T 9F": {"kind": "analog-value", "value": 7.4, "units": "degrees-celsius"},
        "冷冻水温度": {"kind": "analog-value", "value": 8.0, "units": "degrees-celsius"},
        "X" * 80: {"kind": "analog-value", "value": 9.0, "units": "degrees-celsius"},
        " ": {"kind": "binary-value", "value": "active"},
    }


@pytest.fixture()
async def env():
    net_name = f"disc-{uuid.uuid4().hex[:8]}"
    net = VirtualNetwork(net_name)
    dirty = DirtySimDevice(
        net,
        net_name,
        device_id=519,
        mac=2,
        rpm_supported=True,
        objects=DirtySimDevice.DIRTY_OBJECTS,
    )
    old = SimDevice(net, net_name, device_id=620, mac=3, rpm_supported=False)
    await dirty.apply_spec(DirtySimDevice.DIRTY_OBJECTS)
    await old.apply_spec()
    app = build_virtual_application(net, net_name, mac=1, device_id=900)
    stack = BacnetStack(app, max_inflight=2, inter_request_gap_ms=0)
    return stack, dirty, old, net


class TestDiscover:
    async def test_three_phases_and_marks(self, env, tmp_path):
        stack, _dirty, _old, _ = env
        report = await Discoverer(stack).run()
        # Phase 1：两台设备入账
        devs = {d.device_id: d for d in report.devices}
        assert set(devs) == {519, 620}
        assert devs[519].address == "2"
        assert devs[620].object_count > 0  # 老设备（RPM 不支持）也完成枚举

        by_original = {p.original_name: p for p in report.points if p.device_id == 519}
        # §3.3 分支逐个触发
        assert "sanitized" in by_original["CHWS T 9F"].marks
        assert by_original["CHWS T 9F"].raw_name == "CHWS_T_9F"
        assert "sanitized" in by_original["冷冻水温度"].marks
        long_one = next(
            p for p in report.points if p.device_id == 519 and len(p.original_name.strip()) == 80
        )
        assert "truncated" in long_one.marks and len(long_one.raw_name) == 64
        generated = [
            p for p in report.points if p.device_id == 519 and p.obj_type == "binary-value"
        ]
        assert generated and "generated" in generated[0].marks

        # §4.3 单位映射：degrees-celsius → degC
        assert by_original["CHWS T 9F"].unit_token == "degC"

    async def test_products_report_staging_xlsx(self, env, tmp_path):
        stack, _dirty, _old, _ = env
        report = await Discoverer(stack).run()
        report_path = write_report(report, str(tmp_path))
        assert report_path.exists()
        data = json.loads(report_path.read_text(encoding="utf-8"))
        assert data["points"] and data["devices"]

        db = open_db(str(tmp_path / "cache.db"))
        staged = stage_pointmap(db, report)
        assert staged == len(report.points)

        # xlsx：M2 §5.1 canonical 四列，列头逐字
        xlsx = write_export_xlsx(report, str(tmp_path))
        assert xlsx.exists()
        from openpyxl import load_workbook

        wb = load_workbook(xlsx)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        assert header == ["点号", "描述", "单位", "方向"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == len(report.points)
        for _name, desc, _unit, direction in rows:
            assert direction == "" or direction is None  # D 留空（方向从严）
            assert desc.startswith("[")  # 描述带对象类型提示
        adopt = adopt_staging(db)
        assert len(adopt["added"]) == len(report.points)
        db.close()

    async def test_rediscovery_address_drift_detected(self, env, tmp_path):
        """§14.3-3：同名对象地址漂移——报告标记 + adopt 跟随物理地址。"""
        stack, dirty, _old, _net = env
        report1 = await Discoverer(stack).run()
        db = open_db(str(tmp_path / "cache.db"))
        stage_pointmap(db, report1)
        adopt_staging(db)

        # 物理世界变化：519 号设备的 CHWS_T_9F（analog-value:1）删除，
        # 新增一个同名 analog-value:99（地址漂移）
        dirty.app.delete_object(dirty.objects["CHWS T 9F"])
        drift_obj = AnalogValueObject(objectIdentifier=("analog-value", 99), objectName="CHWS T 9F")
        await drift_obj.write_property("presentValue", 6.6)
        dirty.app.add_object(drift_obj)

        report2 = await Discoverer(stack).run()
        stage_pointmap(db, report2)
        adopt = adopt_staging(db)
        drifts = {d["raw_name"]: d for d in adopt["address_drift"]}
        assert "CHWS_T_9F" in drifts
        assert drifts["CHWS_T_9F"]["old"] == [519, "analog-value", 1]
        assert drifts["CHWS_T_9F"]["new"] == [519, "analog-value", 99]
        # raw_name 稳定（云端零感知），地址已跟随物理真相
        row = {p.raw_name: p for p in load_pointmap(db)}["CHWS_T_9F"]
        assert (row.obj_type, row.obj_instance) == ("analog-value", 99)
        db.close()
